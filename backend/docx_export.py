"""Render a stored resume version as a formatted Word document.

Resumes are kept as plain text so the assistant can rewrite them freely. This
module infers the structure back out of that text and lays it out.

Deliberately ATS-safe: no tables, text boxes, headers/footers, images, or
columns. Applicant tracking systems parse those badly or drop them entirely,
and a resume that renders beautifully but parses to mush costs interviews.
Structure comes from real paragraph styles and a genuine bullet list, never
from literal "•" characters or manual spacing.
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "Calibri"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
RULE = "999999"

BULLET_PREFIXES = ("- ", "* ", "• ", "– ", "— ")

# A4: the default for the Indian market this tracker is aimed at.
PAGE_WIDTH_MM = 210
PAGE_HEIGHT_MM = 297


def _is_bullet(line: str) -> bool:
    return line.startswith(BULLET_PREFIXES)


def _strip_bullet(line: str) -> str:
    for prefix in BULLET_PREFIXES:
        if line.startswith(prefix):
            return line[len(prefix) :].strip()
    return line.strip()


def _is_heading(line: str) -> bool:
    """A section heading: short, all-caps, no sentence punctuation."""
    stripped = line.strip()
    if not stripped or len(stripped) > 60 or _is_bullet(stripped):
        return False
    letters = [c for c in stripped if c.isalpha()]
    if not letters:
        return False
    return all(c.isupper() for c in letters) and not stripped.endswith((".", ":", ","))


def parse(text: str) -> dict[str, Any]:
    """Split resume text into a name, contact lines, and titled sections."""
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    name = ""
    contact: list[str] = []
    sections: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    # The first line is always the name, even though it's usually written in
    # caps and would otherwise look exactly like a section heading.
    if lines:
        name = lines.pop(0)

    for line in lines:
        if _is_heading(line):
            current = {"title": line.title(), "blocks": []}
            sections.append(current)
            continue
        if current is None:
            # Everything between the name and the first heading is contact info.
            contact.append(line)
            continue
        current["blocks"].append(
            {"kind": "bullet" if _is_bullet(line) else "line", "text": _strip_bullet(line)}
        )

    # Inside a section that uses bullets, a plain line is a role/heading line.
    for section in sections:
        has_bullets = any(b["kind"] == "bullet" for b in section["blocks"])
        for block in section["blocks"]:
            if has_bullets and block["kind"] == "line":
                block["kind"] = "role"

    return {"name": name, "contact": contact, "sections": sections}


def _bottom_border(paragraph) -> None:
    """Hairline rule under a section heading (python-docx has no API for this)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.makeelement(qn("w:pBdr"), {})
    bottom = borders.makeelement(
        qn("w:bottom"),
        {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "2", qn("w:color"): RULE},
    )
    borders.append(bottom)
    p_pr.append(borders)


def _style(run, size: float, *, bold: bool = False, color: RGBColor = INK, caps: bool = False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if caps:
        run.font.all_caps = True
    return run


def _configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Pt(PAGE_WIDTH_MM * 72 / 25.4)
    section.page_height = Pt(PAGE_HEIGHT_MM * 72 / 25.4)
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Pt(40))
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Pt(46))

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    # East-Asian font mapping, or Word may substitute for some glyphs.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08


def build(resume: dict[str, Any]) -> bytes:
    """Render one resume record to .docx bytes."""
    parsed = parse(resume.get("content", ""))
    document = Document()
    _configure(document)

    name = parsed["name"] or resume.get("name") or "Resume"
    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(2)
    _style(heading.add_run(name), 20, bold=True)

    if not parsed["sections"]:
        # No recognisable structure. parse() has already collected every line
        # after the first into `contact`, so emit those as plain body text
        # rather than styling them as a letterhead.
        for line in parsed["contact"]:
            para = document.add_paragraph()
            para.paragraph_format.space_after = Pt(4)
            _style(para.add_run(line), 10.5)
        return _to_bytes(document)

    for line in parsed["contact"]:
        para = document.add_paragraph()
        para.paragraph_format.space_after = Pt(1)
        _style(para.add_run(line), 10, color=MUTED)

    for section in parsed["sections"]:
        title = document.add_paragraph()
        title.paragraph_format.space_before = Pt(14)
        title.paragraph_format.space_after = Pt(5)
        _style(title.add_run(section["title"]), 11.5, bold=True, caps=True)
        _bottom_border(title)

        for block in section["blocks"]:
            if block["kind"] == "bullet":
                para = document.add_paragraph(style="List Bullet")
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.left_indent = Pt(16)
                _style(para.add_run(block["text"]), 10.5)
            elif block["kind"] == "role":
                para = document.add_paragraph()
                para.paragraph_format.space_before = Pt(7)
                para.paragraph_format.space_after = Pt(3)
                _style(para.add_run(block["text"]), 11, bold=True)
            else:
                para = document.add_paragraph()
                para.paragraph_format.space_after = Pt(3)
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _style(para.add_run(block["text"]), 10.5)

    return _to_bytes(document)


def _to_bytes(document: Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def filename_for(resume: dict[str, Any], person: str = "") -> str:
    """A tidy, filesystem-safe download name."""
    parsed = parse(resume.get("content", ""))
    who = person or parsed["name"] or "Resume"
    label = resume.get("version_label") or resume.get("name") or ""
    stem = f"{who} {label}".strip() or "Resume"
    stem = re.sub(r"[^\w\s-]", "", stem).strip()
    stem = re.sub(r"[\s_-]+", "_", stem)
    return f"{stem[:80]}.docx"
