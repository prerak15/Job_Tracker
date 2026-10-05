"""End-to-end checks for every backend domain, against a throwaway data dir.

Run it from the repo root:

    .venv/Scripts/python.exe backend/selftest.py

No test framework and no server needed — plain asserts, so a fresh clone can
verify itself before touching real data. `jsonstore.DATA_DIR` is repointed at a
temp directory *before* the domain modules are imported, so your own
`data/*.json` is never read or written by this file.
"""

from __future__ import annotations

import asyncio
import json
import sys
import tempfile
from pathlib import Path

BACKEND = Path(__file__).resolve().parent
sys.path.insert(0, str(BACKEND))

import jsonstore  # noqa: E402

tmp = Path(tempfile.mkdtemp(prefix="tracker-selftest-"))
jsonstore.DATA_DIR = tmp
print(f"scratch data dir: {tmp}\n")

import design  # noqa: E402
import dsa  # noqa: E402
import prep  # noqa: E402
import resumes  # noqa: E402
import storage  # noqa: E402

# ------------------------------------------------------------------ jobs
job = storage.create_job(
    {
        "job_title": "Backend Engineer",
        "organisation": "Acme Corp",
        "status": "applied",
        "date_job_posted": "2026-07-20",
        "company_type": "startup",
        "industry": "fintech",
        "org_summary": "Payment infra for SMBs.",
        "source": "linkedin",
        "found_via": "post by their VP Eng",
        "job_description": "Python, Kafka, gRPC, distributed systems.",
    }
)
assert job["id"] and job["date_created"], job
assert job["latest_update"] == "Record created"
assert job["found_via"].startswith("post by"), job["found_via"]

storage.add_contact(
    job["id"],
    {"name": "Jane Doe", "role": "recruiter", "is_referral": True, "email": "j@acme.com"},
)
assert storage.get_job(job["id"])["referred_by"] == "Jane Doe"

storage.add_round(
    job["id"], {"name": "Online assessment", "result": "cleared", "feedback": "DP was hard"}
)
job = storage.get_job(job["id"])
assert job["status"] == "interviewing", job["status"]
assert job["rounds"][0]["round"] == 1

storage.update_round(job["id"], 1, {"feedback": "DP was hard; revise LIS"})
storage.add_followup(job["id"], "Emailed Jane")
job = storage.get_job(job["id"])
assert len(job["follow_ups_sent"]) == 1
assert job["rounds"][0]["feedback"].endswith("revise LIS")

job2 = storage.create_job(
    {"job_title": "SDE", "organisation": "Globex", "status": "rejected"}
)

s = storage.stats()
assert s["total"] == 2 and s["applied_total"] == 2, s
assert s["interview_rate"] == 50.0, s["interview_rate"]
assert s["response_rate"] == 100.0, s["response_rate"]
assert s["rejection_rate"] == 50.0, s["rejection_rate"]
assert s["by_company_type"]["startup"]["applications"] == 1
print("jobs OK    ", {k: s[k] for k in ("total", "response_rate", "interview_rate")})

# A freshly-created record must not mask a deliberately older update date, or
# nothing ever looks overdue.
storage.update_job(job2["id"], {"status": "applied", "latest_update_date": "2020-01-01"})
due = storage.followup_suggestions(stale_days=7)
assert len(due) == 1 and due[0]["organisation"] == "Globex", due
assert due[0]["days_since_activity"] > 1000
print("followups OK", due[0]["days_since_activity"], "days stale")

# ------------------------------------------------------------------ resumes
master = resumes.create_resume(
    {
        "name": "SDE master",
        "version_label": "v1",
        "is_master": True,
        "content": "Python, Django, Postgres",
    }
)
resumes.link_to_job(master["id"], job["id"])
assert storage.get_job(job["id"])["resume_id"] == master["id"]

resumes.add_tailoring(
    master["id"],
    {
        "job_id": job["id"],
        "missing_keywords": ["Kafka", "gRPC"],
        "suggestions": [{"section": "Experience", "change": "Reframe around event-driven"}],
    },
)
r = resumes.get_resume(master["id"])
assert r["tailoring"][0]["organisation"] == "Acme Corp", r["tailoring"][0]
assert r["tailoring"][0]["applied"] is False
resumes.set_tailoring_applied(master["id"], 0, True)
assert resumes.get_resume(master["id"])["tailoring"][0]["applied"] is True

# LaTeX-only version: plain text must be derived so the dashboard and the Word
# export still work.
latex_only = resumes.create_resume(
    {
        "name": "LaTeX master",
        "is_latex_template": True,
        "latex_content": (
            "\\documentclass{article}\\begin{document}"
            "\\textbf{Jane Roe} \\\\ jane@example.com"
            "\\section{Education}"
            "\\end{document}"
        ),
    }
)
assert latex_only["content"].strip(), "plain text should be derived from LaTeX"
assert "documentclass" not in latex_only["content"], latex_only["content"]

rs = resumes.stats()
version = next(v for v in rs["by_version"] if v["id"] == master["id"])
assert version["applications"] == 1 and version["interview_rate"] == 100.0, version
print("resumes OK ", {k: version[k] for k in ("applications", "interview_rate")})

# ------------------------------------------------------------------ dsa
p = dsa.create_problem(
    {
        "title": "Longest Increasing Subsequence",
        "difficulty": "medium",
        "topics": ["dp", "binary-search"],
        "status": "in_progress",
        "phase": "phase_4",
    }
)
assert p["date_started"] and p["date_completed"] is None, p
assert p["phase"] == "phase_4" and p["time_complexity"] == "", p
dsa.log_issue(p["id"], "Missed the patience-sorting angle")
p = dsa.update_problem(
    p["id"],
    {
        "status": "solved",
        "time_spent_minutes": 55,
        "attempts": 2,
        "used_hint": True,
        "confidence": 2,
        "time_complexity": r"O(n \log n)",
        "space_complexity": "O(n)",
    },
)
assert p["date_completed"] and len(p["issues"]) == 1
assert p["time_complexity"] == r"O(n \log n)", p["time_complexity"]

two_sum = dsa.create_problem(
    {
        "title": "Two Sum",
        "difficulty": "easy",
        "topics": ["hashmap"],
        "status": "solved",
        "confidence": 5,
        "time_spent_minutes": 10,
        "phase": "phase_1",
    }
)
dsa.create_problem(
    {"title": "Median of Two Sorted Arrays", "difficulty": "hard", "status": "stuck"}
)

ds = dsa.stats()
assert ds["total"] == 3 and ds["solved"] == 2, ds
assert ds["by_difficulty"]["hard"]["solved"] == 0
assert ds["avg_time_minutes"] == 32.5, ds["avg_time_minutes"]
assert ds["hint_rate"] == 50.0, ds["hint_rate"]
assert ds["by_phase"]["phase_4"]["solved"] == 1, ds["by_phase"]
assert ds["by_phase"]["unassigned"]["total"] == 1, ds["by_phase"]

# Two Sum is solved but unanalysed: it is counted, and it must NOT be pushed
# into the revision queue -- that queue means "re-solve this", and an
# unannotated solve needs a one-line note instead.
assert ds["missing_complexity"] == 1, ds["missing_complexity"]
assert [m["id"] for m in dsa.missing_complexity()] == [two_sum["id"]]

q = dsa.revision_queue()
assert len(q) == 1 and q[0]["title"].startswith("Longest"), q
assert "low confidence" in q[0]["reason"], q[0]["reason"]
print("dsa OK     ", {k: ds[k] for k in ("total", "solved", "missing_complexity")})

dsa.log_revisit(p["id"], "solved unaided in 20m", confidence=5)
assert len(dsa.revision_queue()) == 0, dsa.revision_queue()

# ------------------------------------------------------------------ design
t = design.create_topic(
    {
        "title": "Design a URL shortener",
        "kind": "hld",
        "concepts": ["consistent-hashing", "caching"],
        "status": "studying",
    }
)
assert t["date_started"] and t["date_completed"] is None
design.log_issue(t["id"], "Couldn't justify the sharding key")
t = design.update_topic(
    t["id"],
    {"status": "practiced", "time_spent_minutes": 90, "confidence": 3, "tradeoffs": "base62"},
)
assert t["date_completed"], t
design.add_artifact(t["id"], {"type": "diagram", "path_or_url": "design/url.excalidraw"})
design.create_topic(
    {"title": "Parking lot", "kind": "lld", "patterns": ["strategy"], "status": "practiced",
     "confidence": 5}
)

gs = design.stats()
assert gs["total"] == 2 and gs["practiced"] == 2, gs
assert gs["artifacts"] == 1
assert len(design.revision_queue()) == 1
print("design OK  ", {k: gs[k] for k in ("total", "practiced", "completion_rate")})

# ------------------------------------------------------------------ prep
prep.upsert_phase({"key": "phase_1", "name": "Trees", "status": "completed", "order": 1})
prep.upsert_phase({"key": "phase_4", "name": "DP", "status": "current", "order": 4})
prep.update_profile({"target_companies": ["Acme Corp"], "target_levels": ["SWE2"]})

profile = prep.get_prep()
assert [ph["key"] for ph in profile["phases"]] == ["phase_1", "phase_4"], profile["phases"]
# Counts are joined in from dsa by phase key.
assert profile["phases"][1]["problems"] == 1 and profile["phases"][1]["solved"] == 1

# Setting one phase current must demote the other, or current_phase is ambiguous.
prep.set_phase_status("phase_1", "current")
after = prep.get_prep()
assert after["profile"]["current_phase"] == "phase_1", after["profile"]
assert next(p for p in after["phases"] if p["key"] == "phase_4")["status"] == "upcoming"
assert prep.set_phase_status("nope", "current") is None

prep.set_milestone("Milestone 2 - cycle detection")
assert prep.get_prep()["profile"]["current_milestone"].startswith("Milestone 2")

habit = prep.add_standing_issue("Forgets to return the recursive call", "logic")
prep.flag_standing_issue(habit["id"], problem_id=p["id"])
prep.flag_standing_issue(habit["id"], problem_id=two_sum["id"])
typo = prep.add_standing_issue("Uses [r, c] in a set", "syntax")
assert prep.list_standing_issues()[0]["id"] == habit["id"], "most recurrent should sort first"
assert len(prep.list_standing_issues(active_only=True)) == 2

prep.resolve_standing_issue(typo["id"])
resolved = next(i for i in prep.list_standing_issues() if i["id"] == typo["id"])
assert resolved["active"] is False and resolved["date_resolved"], resolved
assert len(prep.list_standing_issues(active_only=True)) == 1

# Flagging a resolved habit reopens it -- the whole point of tracking recurrence.
prep.flag_standing_issue(typo["id"])
assert prep.list_standing_issues(active_only=True).__len__() == 2

ps = prep.stats()
assert ps["current_phase"] == "Trees", ps["current_phase"]
assert ps["phases_completed"] == 0 and ps["phases_total"] == 2, ps
assert ps["standing_issues_active"] == 2, ps
assert ps["recurring"][0]["times_seen"] == 2, ps["recurring"]
assert prep.delete_standing_issue(typo["id"]) is True
assert prep.delete_standing_issue("nope") is False
print("prep OK    ", {k: ps[k] for k in ("current_phase", "standing_issues_active")})

# ------------------------------------------------------------------ readiness
storage.add_round(job["id"], {"name": "Onsite", "date": "2099-01-01", "result": "pending"})
ready = prep.readiness()
assert len(ready["upcoming_rounds"]) == 1, ready["upcoming_rounds"]
assert ready["upcoming_rounds"][0]["organisation"] == "Acme Corp"
assert ready["next_round_in_days"] > 0
# Acme is live but nothing is tagged to it, which is exactly the gap to surface.
assert "Acme Corp" in ready["untagged_companies"], ready
assert [x["title"] for x in ready["dsa_stuck"]] == ["Median of Two Sorted Arrays"]

dsa.update_problem(two_sum["id"], {"company_tags": ["acme corp"]})
ready = prep.readiness()
coverage = next(c for c in ready["company_coverage"] if c["organisation"] == "Acme Corp")
assert coverage["dsa_tagged"] == 1 and coverage["dsa_solved"] == 1, coverage
assert "Acme Corp" not in ready["untagged_companies"]
print("readiness OK", {k: ready[k] for k in ("next_round_in_days", "dsa_missing_complexity")})

# ------------------------------------------------------------------ next up
# phase_1 ("Trees") is current at this point. Two fresh todos: one inside the
# current phase, one with no phase at all.
in_phase = dsa.create_problem(
    {"title": "Validate BST", "difficulty": "medium", "status": "todo", "phase": "phase_1"}
)
backlog = dsa.create_problem({"title": "Two Pointers Warmup", "difficulty": "easy"})

nxt = dsa.next_up()
assert [n["title"] for n in nxt] == [
    "Median of Two Sorted Arrays",  # stuck -- unblock before starting anything
    "Validate BST",  # todo, current phase
    "Two Pointers Warmup",  # todo, no phase, so it sorts below despite being easy
], [n["title"] for n in nxt]
assert "stuck" in nxt[0]["reason"], nxt[0]["reason"]
assert nxt[1]["in_current_phase"] and "Trees" in nxt[1]["reason"], nxt[1]
assert nxt[2]["in_current_phase"] is False, nxt[2]

# The two queues must stay disjoint: next_up() is "attempt this", the revision
# queue is "re-solve this". A solved problem can never appear here.
assert all(n["status"] in dsa.NEXT_UP_RANK for n in nxt), nxt
assert not {n["id"] for n in nxt} & {r["id"] for r in dsa.revision_queue()}

# Open work outranks the plan -- finish what is started before starting more.
dsa.update_problem(backlog["id"], {"status": "in_progress"})
assert dsa.next_up()[0]["title"] == "Two Pointers Warmup", dsa.next_up()
assert dsa.next_up()[0]["days_open"] == 0, dsa.next_up()[0]
assert len(dsa.next_up(limit=1)) == 1
dsa.delete_problem(in_phase["id"])
dsa.delete_problem(backlog["id"])
print("next up OK ", [n["title"] for n in nxt])

# ------------------------------------------------------------------ coach
# Nothing open, no entrenched habit yet, nothing due for revision: the coach
# falls through to the curriculum and behaves like next_up().
c = dsa.coach()
assert c["kind"] == "advance" and c["pick"]["title"].startswith("Median"), c

# Deferring pulls it off every queue but leaves a checkable way back in.
median_id = next(p["id"] for p in dsa.list_problems() if p["title"].startswith("Median"))
gate = dsa.create_problem({"title": "In-degree warmup", "difficulty": "easy", "phase": "phase_1"})
held = dsa.defer_problem(median_id, "too far ahead of the fundamentals", [gate["id"]])
assert held["status"] == "deferred" and held["defer"]["until_solved"] == [gate["id"]], held
assert not any(n["title"].startswith("Median") for n in dsa.next_up()), dsa.next_up()

c = dsa.coach()
assert [h["title"] for h in c["on_hold"]] == ["Median of Two Sorted Arrays"], c["on_hold"]
assert c["on_hold"][0]["blockers"] == ["In-degree warmup"], c["on_hold"]
assert c["kind"] == "advance" and c["pick"]["title"] == "In-degree warmup", c

# Solving the prerequisite reopens it on its own -- nobody has to remember.
dsa.update_problem(gate["id"], {"status": "solved", "confidence": 5})
c = dsa.coach()
assert c["kind"] == "unlocked" and c["pick"]["title"].startswith("Median"), c
assert c["on_hold"] == [] and len(c["unlocked"]) == 1, c

# Leaving deferred must clear the gate, or a live problem still reads as held.
assert dsa.update_problem(median_id, {"status": "todo"})["defer"] is None

# Open work outranks the plan.
dsa.update_problem(gate["id"], {"status": "in_progress"})
assert dsa.coach()["kind"] == "finish", dsa.coach()
dsa.update_problem(gate["id"], {"status": "solved"})

# A habit seen three times outranks the curriculum -- and the drill is a
# *different* problem on the same concept, because reproducing an answer you
# have already seen tests recall, not whether the idea transferred.
prep.flag_standing_issue(habit["id"], problem_id=two_sum["id"])
assert len(prep.list_standing_issues()[0]["seen_on"]) == 3
same_concept = dsa.create_problem(
    {"title": "Contains Duplicate", "difficulty": "easy", "topics": ["hashmap"]}
)
c = dsa.coach()
assert c["kind"] == "drill" and c["pick"]["action"] == "solve", c
assert c["pick"]["title"] == "Contains Duplicate", c["pick"]
assert any("recurred 3x" in b for b in c["because"]), c["because"]

# With nothing unsolved on that concept, the drill falls back to a redo.
dsa.delete_problem(same_concept["id"])
c = dsa.coach()
assert c["kind"] == "drill" and c["pick"]["action"] == "redo", c
assert c["pick"]["title"] == "Two Sum", c["pick"]

# Habit beaten, but three solves since the last revisit: revision cadence wins.
prep.resolve_standing_issue(habit["id"])
dsa.update_problem(two_sum["id"], {"confidence": 2})
c = dsa.coach()
assert c["kind"] == "revise" and c["pick"]["action"] == "redo", c
assert c["solved_since_last_revisit"] >= dsa.REVISE_EVERY, c
print("coach OK   ", {k: c[k] for k in ("kind", "revision_due")})

# ------------------------------------------------------------------ stopwatch
import datetime as _dt  # noqa: E402

timed = dsa.create_problem({"title": "Timed problem", "difficulty": "easy"})
assert timed["elapsed_seconds"] == 0 and timed["timer"]["started_at"] is None, timed

# Start is the clock *and* the status change -- two clicks is how a timer ends
# up never being used.
running = dsa.set_timer(timed["id"], "start")
assert running["status"] == "in_progress" and running["date_started"], running
assert running["timer"]["started_at"], running["timer"]
assert dsa.stats()["timer_running"]["id"] == timed["id"], dsa.stats()["timer_running"]

# Backdate the running segment rather than sleeping -- the selftest stays fast.
def _rewind(problem_id, seconds):
    data = dsa.load()
    for record in data[dsa.KEY]:
        if record["id"] == problem_id and record["timer"]["started_at"]:
            record["timer"]["started_at"] = (
                _dt.datetime.now() - _dt.timedelta(seconds=seconds)
            ).isoformat(timespec="seconds")
    dsa.write(dsa.FILE, data)

_rewind(timed["id"], 300)
assert 299 <= dsa.get_problem(timed["id"])["elapsed_seconds"] <= 302, dsa.get_problem(timed["id"])

paused = dsa.set_timer(timed["id"], "pause")
assert paused["timer"]["started_at"] is None, paused["timer"]
assert 299 <= paused["timer"]["accumulated_seconds"] <= 302, paused["timer"]
assert dsa.stats()["timer_running"] is None

# Paused time is not counted, and resuming adds to the total rather than
# restarting it -- that is the whole point of having a pause.
dsa.set_timer(timed["id"], "start")
_rewind(timed["id"], 120)
banked = dsa.update_problem(timed["id"], {"status": "solved"})
assert banked["timer"]["started_at"] is None, banked["timer"]
assert 419 <= banked["timer"]["accumulated_seconds"] <= 423, banked["timer"]
assert banked["time_spent_minutes"] == 7, banked["time_spent_minutes"]

# A hand-entered figure is an explicit act and must survive the roll-up.
manual = dsa.create_problem({"title": "Solved offline", "time_spent_minutes": 45})
dsa.set_timer(manual["id"], "start")
_rewind(manual["id"], 90)
manual = dsa.update_problem(manual["id"], {"status": "solved"})
assert manual["time_spent_minutes"] == 45, manual["time_spent_minutes"]

# A timer left running overnight would wreck avg_time_minutes on its own, so a
# single segment is banked at the cap and flagged rather than silently trusted.
forgotten = dsa.create_problem({"title": "Left running"})
dsa.set_timer(forgotten["id"], "start")
_rewind(forgotten["id"], (dsa.STALE_SEGMENT_HOURS + 9) * 3600)
forgotten = dsa.set_timer(forgotten["id"], "pause")
assert forgotten["timer"]["capped"] is True, forgotten["timer"]
assert forgotten["timer"]["accumulated_seconds"] == dsa.STALE_SEGMENT_HOURS * 3600, forgotten

assert dsa.set_timer(forgotten["id"], "reset")["elapsed_seconds"] == 0
try:
    dsa.set_timer(forgotten["id"], "rewind")
    raise AssertionError("an unknown timer action must not be accepted")
except ValueError:
    pass

# The attempt cap. Under the line nothing happens; over it -- running or paused
# for review -- the attempt is owed its solution.
grind = dsa.create_problem({"title": "Long grind"})
dsa.set_timer(grind["id"], "start")
_rewind(grind["id"], (dsa.ATTEMPT_CAP_MINUTES - 1) * 60)
assert dsa.get_problem(grind["id"])["over_attempt_cap"] is False
assert dsa.over_attempt_cap() == []
_rewind(grind["id"], (dsa.ATTEMPT_CAP_MINUTES + 3) * 60)
dsa.set_timer(grind["id"], "pause")
assert dsa.get_problem(grind["id"])["over_attempt_cap"] is True, "a paused clock still crossed the line"
assert [p["id"] for p in dsa.over_attempt_cap()] == [grind["id"]]

capped = dsa.cap_attempt(grind["id"], on="2026-09-22")
assert capped["status"] == "deferred" and capped["defer"]["review_on"] == "2026-10-06", capped["defer"]
assert capped["attempts"] == 1 and capped["used_hint"] is True, capped
assert f"after {dsa.ATTEMPT_CAP_MINUTES + 3} min" in capped["issues"][-1]["issue"], capped["issues"]
# Zeroed, because the cap is per attempt: the return visit must start fresh
# rather than already over the line.
assert capped["timer"]["accumulated_seconds"] == 0 and capped["elapsed_seconds"] == 0
assert capped["over_attempt_cap"] is False and dsa.over_attempt_cap() == []
assert dsa.cap_attempt("no-such-id") is None

# elapsed_seconds is derived, so it must never reach the file.
raw = json.loads((tmp / "dsa.json").read_text(encoding="utf-8"))
assert all("elapsed_seconds" not in p for p in raw["problems"]), "derived field was persisted"
assert all("over_attempt_cap" not in p for p in raw["problems"]), "derived field was persisted"
print("stopwatch OK", {k: dsa.stats()[k] for k in ("avg_time_minutes", "timed")})

# ------------------------------------------------------------- patterns
import pattern_seed  # noqa: E402
import patterns  # noqa: E402

# The join key is namespaced by host: LeetCode and GfG both publish
# /problems/<slug>, and unnamespaced they would claim each other's problems.
assert patterns.slug_of("https://leetcode.com/problems/two-sum/description/") == "leetcode:two-sum"
assert patterns.slug_of("https://www.geeksforgeeks.org/problems/two-sum/1") == "geeksforgeeks:two-sum"
assert patterns.slug_of("https://leetcode.com/problems/two-sum/") != patterns.slug_of(
    "https://www.geeksforgeeks.org/problems/two-sum/1"
), "two sites must not share a slug"
assert patterns.slug_of("https://example.com/blog/two-sum") is None
# His titles carry the LeetCode number; the sheet's do not.
assert patterns.norm_title("LC 207 - Course Schedule") == patterns.norm_title("Course Schedule")
assert patterns.norm_title("leetcode 200: Number of Islands") == "numberofislands"

patterns.seed(
    [
        {
            "key": "sliding_window",
            "name": "Sliding Window",
            "domain": "dsa",
            "order": 1,
            "idea": "Both ends only move forward.",
            "problems": [
                {
                    "title": "Longest Substring Without Repeating Characters",
                    "url": "https://leetcode.com/problems/longest-substring-without-repeating-characters/",
                    "difficulty": "medium",
                },
                {"title": "Two Sum", "url": None},
                {"title": "Minimum Window Substring", "url": None, "difficulty": "hard"},
            ],
        },
        {
            "key": "hld_caching",
            "name": "Caching",
            "domain": "hld",
            "order": 2,
            "idea": "Every cache design is an invalidation design.",
            "problems": [{"title": "Design a news feed", "url": None}],
        },
    ]
)
again = patterns.seed(pattern_seed.PATTERNS)
assert again["skipped"] == 2 and again["added"] == len(pattern_seed.PATTERNS) - 2, again
assert patterns.seed(pattern_seed.PATTERNS)["added"] == 0, "seeding must be idempotent"

window = patterns.get_pattern("sliding_window")
# "Two Sum" has no URL in this catalogue, so it can only match by title -- and
# there is a solved Two Sum from the dsa section above.
by_title = next(q for q in window["problems"] if q["title"] == "Two Sum")
assert by_title["tracked"] and by_title["solved"], by_title
assert window["solved"] == 1 and window["total"] == 3, window
assert window["state"] == "learning" and window["coverage"] == 33.3, window

# A fresh solve joins by slug across two differently-worded titles.
lss = dsa.create_problem(
    {
        "title": "LC 3 - Longest Substring Without Repeating Characters",
        "url": "https://leetcode.com/problems/longest-substring-without-repeating-characters/description/",
        "status": "solved",
        "topics": ["sliding-window"],
        "confidence": 5,
    }
)
window = patterns.get_pattern("sliding_window")
assert window["solved"] == 2 and window["coverage"] == 66.7, window
assert window["state"] == "practiced", window["state"]
assert patterns.for_problem(lss["url"]) == [
    {"key": "sliding_window", "name": "Sliding Window", "domain": "dsa"}
], patterns.for_problem(lss["url"])

# Covered, never rated and never revisited -- that is the one due reason a
# freshly finished pattern can have, and re-deriving it is what clears it.
assert window["due"] and window["due_reason"] == "covered but never rated or revisited", window
patterns.log_revisit("sliding_window", "re-derived the shrink condition", confidence=4)
window = patterns.get_pattern("sliding_window")
assert not window["due"] and window["confidence"] == 4 and window["last_revised"], window

# A flag forces it back; a revisit is what takes it out again.
patterns.update_pattern("sliding_window", {"flagged": True})
assert patterns.get_pattern("sliding_window")["due_reason"] == "flagged for revision"
patterns.log_revisit("sliding_window", "clean second pass", confidence=5)
after = patterns.get_pattern("sliding_window")
assert not after["flagged"] and not after["due"], after

# Nothing untouched is ever due, a hand-set flag included -- that is what keeps
# the two queues disjoint by construction rather than by convention.
patterns.update_pattern("kadane", {"flagged": True})
assert patterns.get_pattern("kadane")["due"] is False, "an untouched pattern must not be due"
due_keys = {q["key"] for q in patterns.revision_queue()}
new_keys = {q["key"] for q in patterns.unstarted()}
assert not (due_keys & new_keys), due_keys & new_keys

# Solved work tagged with a pattern but not listed under it counts as activity,
# and deliberately NOT toward coverage -- a denominator that grows on tagging
# is a percentage that means nothing.
dsa.create_problem(
    {"title": "Fruit Into Baskets", "status": "solved", "topics": ["sliding-window"]}
)
after = patterns.get_pattern("sliding_window")
assert after["extra_solved"] == 1 and after["total"] == 3, after
assert after["coverage"] == 66.7, "off-catalogue work must not move coverage"

# Promotion queues a catalogue row as real work, tagged so it joins back.
promoted = patterns.promote("sliding_window", 2)
assert promoted["domain"] == "dsa"
assert promoted["record"]["status"] == "todo" and promoted["record"]["topics"] == ["sliding-window"]
assert promoted["record"]["difficulty"] == "hard", "the sheet's difficulty carries through"
assert promoted["record"]["phase"] is None, "pattern study must not claim a curriculum phase"
assert patterns.get_pattern("sliding_window")["problems"][2]["tracked"]
try:
    patterns.promote("sliding_window", 2)
    raise AssertionError("promoting an already-tracked problem must fail")
except ValueError:
    pass

design_side = patterns.promote("hld_caching", 0)
assert design_side["domain"] == "hld" and design_side["record"]["kind"] == "hld"
cache = patterns.get_pattern("hld_caching")
assert cache["tracked"] == 1 and cache["solved"] == 0, cache
design.update_topic(design_side["record"]["id"], {"status": "practiced"})
cache = patterns.get_pattern("hld_caching")
assert cache["solved"] == 1 and cache["state"] == "practiced", cache

try:
    patterns.add_problem("sliding_window", {"title": "two sum"})
    raise AssertionError("a duplicate catalogue title must be rejected")
except ValueError:
    pass

# Every progress number is joined at read time; none of it may be persisted.
raw = json.loads((tmp / "patterns.json").read_text(encoding="utf-8"))
derived = {"solved", "total", "coverage", "state", "due", "tracked", "extra_solved"}
assert all(not derived & set(p) for p in raw["patterns"]), "a derived field was persisted"
assert all("id" not in q for p in raw["patterns"] for q in p["problems"]), "catalogue rows are not records"

# The shipped curriculum itself: keys unique, order dense, every idea written.
keys = [p["key"] for p in pattern_seed.PATTERNS]
assert len(set(keys)) == len(keys), "duplicate pattern key in the seed"
assert [p["order"] for p in pattern_seed.PATTERNS] == list(range(1, len(keys) + 1))
for entry in pattern_seed.PATTERNS:
    assert entry["domain"] in patterns.DOMAINS, entry["key"]
    assert entry["idea"].strip(), f"{entry['key']} has no idea written"
    assert entry["problems"], f"{entry['key']} has no problems"

ps = patterns.stats()
assert ps["total"] == len(pattern_seed.PATTERNS), ps
assert ps["due"] + ps["unstarted"] <= ps["total"], ps
print(
    f"patterns OK {ps['total']} patterns, {ps['problems']} problems, "
    f"{ps['due']} due, {ps['unstarted']} unstarted"
)

# ------------------------------------------------------------------ skills
import skill_seed  # noqa: E402
import skills  # noqa: E402

# Matching is the whole domain: a careless alias invents demand that was never
# there, and nothing about the number it produces looks wrong. Every case here
# is one that actually fired against data/jobs.json while the board was built.
assert skills.mentions("Kubernetes at scale", ["Kubernetes"]) == "Kubernetes"
assert skills.mentions("features live within weeks", ["EKS"]) is None
assert skills.mentions("building scalable solutions", ["Scala"]) is None
assert skills.mentions("Java, Ruby, Clojure, Scala, C", ["Scala"]) == "Scala"
assert skills.mentions("JavaScript and TypeScript", ["Java"]) is None
assert skills.mentions("Retail, Sales, Operations", ["AI"]) is None
assert skills.mentions("experience with C++ and C", ["C++"]) == "C++"
# Plurals: postings are written in them, and an alias in the singular saw none
# of the Weekday voice role at all.
assert skills.mentions("authored design documents", ["design document"])
assert skills.mentions("maintain voice pipelines", [skill_seed._VOICE])
assert skills.mentions("autonomous voice agent", [skill_seed._VOICE])
# The guarded ones. Each exists because of a specific false positive.
assert skills.mentions("Go ahead, apply anyway", [skill_seed._GO]) is None
assert skills.mentions("Working knowledge of Go and Spark", [skill_seed._GO]) == "Go"
assert skills.mentions("Java/Go; REST APIs", [skill_seed._GO]) == "Go"
assert skills.mentions("Redis/BullMQ Queues, SSE Streaming", [skill_seed._STREAMING]) is None
assert skills.mentions("batch and streaming data", [skill_seed._STREAMING])

# A skill with explicit aliases must not silently gain its name as one: the Go
# row lists a guard precisely to keep "Go ahead" out, and appending the bare
# name behind it put the false positive straight back.
guarded = skills._normalize({"key": "golang", "name": "Go", "aliases": [skill_seed._GO]})
assert guarded["aliases"] == [skill_seed._GO], guarded["aliases"]
assert skills._normalize({"key": "k", "name": "Kafka"})["aliases"] == ["Kafka"]

# A claim is made in the skills section; prose is evidence, not a claim. This
# resume mentions throughput in a bullet and lists C/C++ under Languages, and
# only the second is a claim about a capability.
resume_text = """Prerak Gupta
Bengaluru, India

WORK EXPERIENCE
- Raised outreach throughput from 30 to 200 emails per hour.

SKILLS & CERTIFICATIONS
Languages: Python, SQL, C/C++
"""
claims = skills.claim_text(resume_text)
assert "C/C++" in claims and "throughput" not in claims, claims
assert skills.claim_text("no headings here") == "no headings here"

seeded = skills.seed(skill_seed.SKILLS)
assert seeded["added"] == len(skill_seed.SKILLS), seeded
assert skills.seed(skill_seed.SKILLS)["added"] == 0, "seeding twice must add nothing"

seed_keys = [s["key"] for s in skill_seed.SKILLS]
assert len(set(seed_keys)) == len(seed_keys), "duplicate skill key in the seed"
for entry in skill_seed.SKILLS:
    assert entry["aliases"], f"{entry['key']} would never match anything"
    assert entry.get("area", "practice") in skills.AREAS, entry["key"]
    assert entry.get("effort", "weeks") in skills.EFFORTS, entry["key"]
    # A row that is neither held nor declined has to say how it gets closed.
    if entry.get("status") not in ("declined", "done") and entry.get("level", 0) < 3:
        assert entry.get("plan", "").strip(), f"{entry['key']} has no plan"

# The job created at the top of this file lists Kafka, gRPC and distributed
# systems, so demand for those is joined in — from jobs.json, on this read.
board = {s["key"]: s for s in skills.list_skills()}
assert board["grpc"]["demand"] == 1, board["grpc"]["mentions"]
assert board["streaming"]["demand"] == 1
assert board["kubernetes"]["demand"] == 0, "nothing in the fixture asks for it"
assert board["grpc"]["organisations"] == ["Acme Corp"]

# Nothing derived may ever reach the file — the same guarantee patterns.py and
# companies.py make. A stored copy of "have I done this" eventually disagrees
# with the domain that owns it, and the stale one is the one being read.
raw = json.loads((tmp / "skills.json").read_text(encoding="utf-8"))["skills"][0]
for derived in ("demand", "mentions", "claimed", "state", "score", "because", "rank"):
    assert derived not in raw, f"{derived} was persisted"

# Demand-free rows stay out of the ranked queue: this board is driven by the
# pipeline, and a row nothing has asked for is not a priority.
queue = skills.gap_queue()
assert all(row["demand"] > 0 for row in queue), "a queue row with no demand"
assert all(row["state"] not in ("have", "declined") for row in queue)
assert [row["rank"] for row in queue] == list(range(1, len(queue) + 1))
assert queue == sorted(queue, key=lambda r: (-r["score"], -r["demand_weighted"], r["name"])), (
    "the queue is not in its own rank order"
)
assert queue[0]["because"], "a ranking nobody can audit gets ignored"

# Declining drops a skill from the queue but not from the demand count — that
# is what lets the decision be re-read later against a number that has moved.
before = len(skills.gap_queue())
skills.decline("grpc", "one posting only")
assert len(skills.gap_queue()) == before - 1
assert skills.get_skill("grpc")["demand"] == 1, "declining must not hide the demand"
assert [d["key"] for d in skills.declined() if d["key"] == "grpc"] == ["grpc"]
try:
    skills.decline("streaming", "   ")
except ValueError:
    pass
else:
    raise AssertionError("declining without a reason must fail")

# Evidence is the only thing that moves a level, and it always leaves a note.
moved = skills.log_evidence("kubernetes", "Deployed the tracker to a kind cluster", 3)
assert moved["level"] == 3 and moved["status"] == "done", moved
assert moved["state"] == "have" and moved["log"][-1]["note"].startswith("Deployed")
assert moved["evidence"].startswith("Deployed")
try:
    skills.log_evidence("kubernetes", "  ")
except ValueError:
    pass
else:
    raise AssertionError("evidence without a note must fail")

# exposed() is allowed to overlap gap_queue() — unlike dsa.py's and patterns.py's
# pairs, they answer different questions and a skill can genuinely need both.
skills.update_skill("java", {"level": 0})
skills.create_skill(
    {
        "key": "selftest-claim",
        "name": "Selftest Claim",
        "aliases": ["Selftest Claim"],
        "level": 0,
        "plan": "n/a",
    }
)
resumes.create_resume(
    {
        "name": "selftest",
        "content": "Name\n\nSKILLS\nLanguages: Selftest Claim\n",
    }
)
claimed = {s["key"]: s for s in skills.list_skills()}["selftest-claim"]
assert claimed["claimed"] and claimed["state"] == "exposed", claimed["claimed_on"]
assert "selftest-claim" in [s["key"] for s in skills.exposed()]

for_job = skills.for_job(job["id"])
assert for_job["has_description"]
assert "grpc" in [s["key"] for s in for_job["declined"]]
assert "python" in [s["key"] for s in for_job["covered"]], for_job["covered"]

sk = skills.stats()
assert sk["total"] == len(skill_seed.SKILLS) + 1, sk
assert sk["jds_read"] >= 1 and 0 <= sk["coverage"] <= 100, sk
assert sk["by_state"]["declined"] >= 1 and sk["exposed"] >= 1
assert len(sk["next"]) <= 3
print(
    f"skills OK   {sk['total']} skills, {sk['jds_read']} JDs read, "
    f"{sk['open']} open, {sk['exposed']} exposed, coverage {sk['coverage']}%"
)

# ------------------------------------------------------- companies + discovery
import companies  # noqa: E402
import company_seed  # noqa: E402
import discover  # noqa: E402

acme = companies.create_company(
    {
        "name": "Acme Corp",
        "category": "startup",
        "tier": "growth",
        "locations": ["Bengaluru"],
        "status": "target",
        "interest": 5,
    }
)
assert acme["applications"] == 1, acme  # the job created at the top of this file
assert acme["board_column"] == "applied", "a live application must beat the manual status"

# Name matching: postings carry legal entities, so a long name matches by
# prefix — but a short one must not, or "Ola" would claim every Olam job.
assert companies.matches(acme, "Acme Corp Private Limited")
assert companies.matches(acme, "acme corp")
short = companies.create_company({"name": "Zap", "aliases": ["Zap Technologies"]})
assert not companies.matches(short, "Zapier Inc"), "short names must match exactly"
assert companies.matches(short, "Zap Technologies Private Limited"), "aliases match by prefix"
assert companies.find_company("zap technologies")["name"] == "Zap"

try:
    companies.create_company({"name": "ACME CORP"})
    raise AssertionError("duplicate company should have been rejected")
except ValueError:
    pass

added = companies.seed(company_seed.SEED)
assert added["added"] == len(company_seed.SEED), added
again = companies.seed(company_seed.SEED)
assert again["added"] == 0 and again["skipped"] == len(company_seed.SEED), again

board = companies.list_companies()
assert len(board) == len(company_seed.SEED) + 2, len(board)
# Sorted by interest descending, so the one thing wanted at 5/5 leads.
assert board[0]["name"] == "Acme Corp", board[0]["name"]
assert all(c["ats_token"] for c in companies.fetchable()), "a wired company needs a token"

# A provider without a token can't be fetched and must not claim it can.
lying = companies.create_company({"name": "No Token Ltd", "ats_provider": "greenhouse"})
assert lying["ats_provider"] == "none", lying

cstats = companies.stats()
assert cstats["total"] == len(board) + 1, cstats["total"]
assert cstats["with_ats"] == len(companies.fetchable())
assert cstats["applied_to"] >= 1
gaps = [g["name"] for g in companies.target_gaps()]
assert "Acme Corp" not in gaps, "a target with an application is not a gap"
companies.update_company(short["id"], {"status": "target", "interest": 4})
assert "Zap" in [g["name"] for g in companies.target_gaps()]
print(f"companies OK {cstats['total']} on board, {cstats['with_ats']} wired")

# --- discovery: pure parsing, no network -----------------------------------
# Shapes mirror the live responses; each provider names things differently and
# only Lever dates in epoch milliseconds.
GREENHOUSE = {
    "jobs": [
        {
            "title": "Software Engineer II",
            "absolute_url": "https://boards.greenhouse.io/acme/jobs/1",
            "location": {"name": "Bengaluru, India"},
            "first_published": "2026-07-20T04:52:12-04:00",
            "updated_at": "2026-07-30T00:00:00-04:00",
            "content": "&lt;p&gt;Build &amp;amp; scale services&lt;/p&gt;",
        },
        {
            "title": "Staff Engineer",
            "absolute_url": "https://boards.greenhouse.io/acme/jobs/2",
            "location": {"name": "San Francisco"},
            "first_published": "2026-07-21T00:00:00-04:00",
        },
    ]
}
LEVER = [
    {
        "text": "Backend Engineer",
        "hostedUrl": "https://jobs.lever.co/acme/abc",
        "categories": {"location": "Bengaluru", "department": "Engineering"},
        "createdAt": 1784539800000,  # 2026-07-20T09:30Z, in epoch milliseconds
        "descriptionPlain": "Go and Kafka.",
    }
]
ASHBY = {
    "jobs": [
        {
            "title": "ML Engineer",
            "jobUrl": "https://jobs.ashbyhq.com/acme/xyz",
            "location": "Bengaluru",
            "secondaryLocations": [{"location": "Chennai"}],
            "publishedAt": "2026-06-29T17:12:35.753+00:00",
            "isListed": True,
            "department": "Research",
        },
        {"title": "Hidden Role", "jobUrl": "https://x/y", "location": "Bengaluru", "isListed": False},
    ]
}

gh = discover.parse("greenhouse", GREENHOUSE, acme)
assert [r["title"] for r in gh] == ["Software Engineer II", "Staff Engineer"]
assert gh[0]["posted"] == "2026-07-20", "first_published wins over updated_at"
assert "Build & scale services" in gh[0]["description"], gh[0]["description"]
lv = discover.parse("lever", LEVER, acme)
assert lv[0]["posted"] == "2026-07-20", lv[0]["posted"]  # epoch ms -> ISO day
ash = discover.parse("ashby", ASHBY, acme)
assert [r["title"] for r in ash] == ["ML Engineer"], "isListed False must be dropped"
assert ash[0]["locations"] == ["Bengaluru", "Chennai"]

# Seniority is word-boundary matched. An earlier version looked for "i " and
# read "Principal AI Security Specialist" as entry level.
for title, want in [
    ("Principal AI Security Specialist", "senior"),
    ("Sr. Staff Software Development Engineer - AI Engineer", "senior"),
    ("Engineering Manager, AI", "senior"),
    ("Software Engineer II", "mid"),
    ("Associate Software Engineer", "entry"),
    ("Associate Director, Platform", "senior"),
    ("Senior Data Science Intern", "intern"),
    # The AI labs title every IC level this way; reading it as senior would
    # hide their whole board from a SWE2 search.
    ("Member of Technical Staff", "mid"),
]:
    assert discover.guess_seniority(title) == want, (title, discover.guess_seniority(title))

# Keywords anchor at the start of a word, so "ai" must not match "Retail".
assert not discover.matches_keywords({"title": "Retail Ops Manager", "department": ""}, ["ai"])
assert discover.matches_keywords({"title": "Engineering Manager", "department": ""}, ["engineer"])

everything = gh + lv + ash
india = discover.filter_roles(everything, keywords=["engineer"], exclude_seniority=["senior"])
# "Staff Engineer" is dropped twice over: senior, and San Francisco.
assert [r["title"] for r in india] == [
    "Software Engineer II",
    "Backend Engineer",
    "ML Engineer",
], india
assert discover.filter_roles(everything, keywords=["sales"]) == []
# Sales titles are dropped by default even when they match an engineering word.
noise = [{**gh[0], "title": "Sales Engineer", "url": "https://x/s"}]
assert discover.filter_roles(noise, keywords=["engineer"]) == []

undated = [{**gh[0], "posted": None}]
assert discover.filter_roles(undated, max_age_days=1), "undated postings must not be hidden"
assert not discover.filter_roles(gh[:1], max_age_days=1, as_of=jsonstore.date.today())

remote = [{**gh[1], "locations": ["Remote"], "remote": True, "seniority": "mid"}]
assert not discover.filter_roles(remote), "remote is opt-in"
assert discover.filter_roles(remote, include_remote=True)

# Dedupe: against what's tracked, by url and by (organisation, title).
# The Lever fixture is "Backend Engineer" at Acme Corp, which is the very job
# created at the top of this file — so a real tracked record must hide it.
kept, already = discover.dedupe(everything, storage.list_jobs())
assert already == 1, already
assert "Backend Engineer" not in [r["title"] for r in kept], kept
tracked = [{"url": gh[0]["url"], "organisation": "", "job_title": ""}]
kept, already = discover.dedupe(everything, tracked)
assert already == 1 and len(kept) == 3, (already, len(kept))
by_title = [{"url": None, "organisation": "Acme Corp", "job_title": "backend engineer"}]
_, already = discover.dedupe(everything, by_title)
assert already == 1, "a hand-added record has no url, so title matching must catch it"
twice = discover.dedupe(everything + everything, [])[0]
assert len(twice) == 4, "the same role listed twice collapses"

promoted = discover.to_job(lv[0])
assert promoted["status"] == "saved", "discovery finds leads, not applications"
assert promoted["source"] == "careers_page"
assert promoted["organisation"] == "Acme Corp" and promoted["date_job_posted"] == "2026-07-20"
print(f"discovery OK {len(everything)} parsed, {len(india)} after filters")

# --- careers-page scraper: pure functions only — no browser, no network ----------
# scrape.py imports Playwright lazily, so this block must pass with it uninstalled.
import roles_xlsx  # noqa: E402
import scrape  # noqa: E402

# URLs: the site's own keyword parameter becomes {q}; a URL without one is untouched.
tpl = scrape.keyword_template("https://x.example/jobs?location=India&q=engineer")
assert "{q}" in tpl and "location=India" in tpl, tpl
assert scrape.fill_url(tpl, "software engineer").count("software+engineer") == 1
plain = "https://x.example/careers?location=India"
assert scrape.keyword_template(plain) == plain
wd = "https://nv.wd5.myworkdayjobs.com/Site?q=x"
assert scrape.default_url_base(wd) == "https://nv.wd5.myworkdayjobs.com/Site"
assert scrape.default_url_base(plain) is None
assert scrape.build_url("https://a/b", page_url="https://p/q", url_base=None) == "https://a/b"
assert scrape.build_url("/job/x", page_url="https://p/q", url_base="https://h/Site") == "https://h/Site/job/x"
assert scrape.build_url("/job/x", page_url="https://p/q/r", url_base=None) == "https://p/job/x"

# JSON mapping: a Workday-shaped payload is found and mapped; a facet list is not jobs.
WORKDAY = {
    "total": 3,
    "facets": [{"id": "IN", "name": "India", "count": 9}] * 3,
    "jobPostings": [
        {"title": "Senior Software Engineer, Platform", "externalPath": "/job/Pune/Sr-SWE_R1",
         "locationsText": "Pune, India", "postedOn": "Posted 3 Days Ago"},
        {"title": "Machine Learning Engineer II", "externalPath": "/job/Bengaluru/MLE_R2",
         "locationsText": "2 Locations", "postedOn": "Posted Today"},
        {"title": "Site Reliability Engineer", "externalPath": "/job/Hyderabad/SRE_R3",
         "locationsText": "Hyderabad, India", "postedOn": "Posted 30+ Days Ago"},
    ],
}
page_two = {"jobPostings": [
    {"title": "Data Platform Engineer", "externalPath": "/job/Pune/DPE_R4",
     "locationsText": "Pune, India", "postedOn": "Posted Yesterday"},
]}
captured = [
    ("https://nv.wd5.myworkdayjobs.com/wday/cxs/nv/Site/jobs?offset=0", WORKDAY),
    ("https://nv.wd5.myworkdayjobs.com/wday/cxs/nv/Site/jobs?offset=20", page_two),
]
base = "https://nv.wd5.myworkdayjobs.com/Site"
found = scrape.find_source(captured, None, None, page_url=base, url_base=base)
assert found, "a job list in a captured response must be found"
endpoint, fmap, rows = found
assert endpoint == "nv.wd5.myworkdayjobs.com/wday/cxs/nv/Site/jobs", endpoint
assert fmap["path"] == "jobPostings" and fmap["title"] == "title" and fmap["url"] == "externalPath", fmap
assert len(rows) == 4, "both pages of the same endpoint merge"
assert rows[0]["url"] == "https://nv.wd5.myworkdayjobs.com/Site/job/Pune/Sr-SWE_R1", rows[0]
assert scrape.guess_field_map(WORKDAY["facets"], "facets") is None, "facets are not roles"
replay = scrape.find_source(captured, endpoint, fmap, page_url=base, url_base=base)
assert replay and len(replay[2]) == 4, "the saved mapping replays"
assert scrape.find_source([("https://x/y", {"a": [1, 2, 3]})], None, None, page_url=base, url_base=None) is None
# Workday reports its total on page one only (later pages say 0): the largest figure wins.
assert scrape.source_total(captured, endpoint, fmap) == 3, "the fixture's own total"
with_total = [(captured[0][0], {**WORKDAY, "total": 244}), (captured[1][0], {**page_two, "total": 0})]
assert scrape.source_total(with_total, endpoint, fmap) == 244
assert [r["url"] for r in scrape.merge_rows(rows[:2], rows[1:])] == [r["url"] for r in rows]

asof = _dt.date(2026, 10, 5)
assert scrape.normalize_posted("Posted 3 Days Ago", asof) == "2026-10-02"
assert scrape.normalize_posted("Posted Today", asof) == "2026-10-05"
assert scrape.normalize_posted("Posted Yesterday", asof) == "2026-10-04"
assert scrape.normalize_posted("2026-09-01T10:00:00Z", asof) == "2026-09-01"
assert scrape.normalize_posted("whenever", asof) is None

# Anchors: job-shaped href and title-shaped text only; location rides in the text.
anchors = [
    {"href": "https://c.example/jobs/12345", "text": "Backend Engineer\nBengaluru, India\nFull time"},
    {"href": "https://c.example/jobs/12345#apply", "text": "Backend Engineer"},
    {"href": "https://c.example/careers", "text": "Careers"},
    {"href": "https://c.example/about", "text": "About our engineering culture"},
    {"href": "https://c.example/job/777", "text": "Apply now"},
    {"href": "/jobs/99", "text": "Platform Engineer"},
]
links_found = scrape.job_links(anchors, "https://c.example/search")
assert [r["title"] for r in links_found] == ["Backend Engineer", "Platform Engineer"], links_found
# Lines two and three both ride along as the location hint: the city is not always second.
assert links_found[0]["location"] == "Bengaluru, India, Full time", links_found[0]
assert links_found[1]["url"] == "https://c.example/jobs/99" and links_found[1]["location"] == ""

# Clusters: a repeated id-bearing link shape is a listing; a block of team pages is not.
teams = ["alpha", "beta", "gamma", "delta", "omega", "sigma"]
listing = [
    {"href": f"https://jobs.x.example/en-in/details/2000{i}-08{i}/software-engineer-{n}?team=T",
     "text": f"Software Engineer {n}\nBengaluru"}
    for i, n in enumerate(teams)
] + [
    {"href": "https://jobs.x.example/en-in/search?location=india", "text": "Search roles here"},
] + [
    {"href": f"https://www.x.example/teams/software-engineering-{n}", "text": f"Software engineering team {n}"}
    for n in teams
]
clustered = scrape.job_links(listing, "https://jobs.x.example/")
assert len(clustered) == 6 and all("/details/" in r["url"] for r in clustered), clustered
assert scrape.cluster_links(listing[7:], "https://x") == [], "team pages carry no ids, so they are not roles"

# An API that exposes only an id: a person supplies the URL pattern once.
ID_ONLY = {"items": [{"roleId": f"R{n}", "jobTitle": f"Platform Engineer {n}", "locations": [{"city": "Pune"}]}
                     for n in range(4)]}
assert scrape.guess_field_map(ID_ONLY["items"], "items") is None, "no link and no template -> not mappable"
templated = {"path": "items", "title": "jobTitle", "id": "roleId", "location": "locations",
             "url_template": "https://h.example/roles/{id}"}
id_rows = scrape.rows_from_source([("https://api.h.example/q", ID_ONLY)], "api.h.example/q", templated,
                                  page_url="https://h.example", url_base=None)
assert [r["url"] for r in id_rows][:2] == ["https://h.example/roles/R0", "https://h.example/roles/R1"]
assert id_rows[0]["location"] == "Pune"
assert scrape.classify_page(202, "  ") and scrape.classify_page(202, "content") is None
# A page that says its list is empty is a real zero, not an unreadable page.
assert scrape.says_no_openings("Open positions\nThere are no job openings currently.")
assert scrape.says_no_openings("We have no open positions right now")
assert scrape.says_no_openings("Currently there are no job postings available."), "Nykaa's wording"
assert scrape.says_no_openings("0 Open jobs available"), "Darwinbox's empty board"
assert not scrape.says_no_openings("10 Open jobs available")
assert scrape.says_no_openings("0\nOpen jobs available"), "the count renders in its own element"

# Darwinbox: a job list with no link field, recognised by its endpoint.
dbx_url = "https://acme.darwinbox.in/ms/candidateapi/job/alljobs?companyId=main"
dbx = {"status": 1, "job_counts": 3, "data": [
    {"id": f"a6ab{i}", "title": t, "locations": "Bangalore, Karnataka, India", "posted_on": 1790620200}
    for i, t in enumerate(["SDE-3 (Backend)", "Staff Engineer", "Customer Success Engineer"])
]}
assert scrape.guess_field_map(dbx["data"], "data") is None, "no link field: the generic guess must not fire"
dbx_found = scrape.find_source([(dbx_url, dbx)], None, None, page_url=dbx_url, url_base=None)
assert dbx_found, "Darwinbox alljobs should be read by its known shape"
dbx_rows = dbx_found[2]
assert dbx_rows[0]["url"] == "https://acme.darwinbox.in/ms/candidatev2/main/careers/jobDetails/a6ab0", dbx_rows[0]
assert dbx_rows[0]["location"] == "Bangalore, Karnataka, India"
assert scrape.source_total([(dbx_url, dbx)], dbx_found[0], dbx_found[1]) == 3, "job_counts is the total"
assert scrape.known_field_map("https://acme.example/ms/candidateapi/job/alljobs", "data") is None
assert scrape.normalize_posted(1790620200) == "2026-09-28", "epoch seconds, not 1970"
assert scrape.normalize_posted(1790620200000) == "2026-09-28", "epoch milliseconds still work"

# Kula: JSON with no link field and the location nested two levels down.
kula_url = "https://careers.kula.ai/api/internal/ats_job_posts?accountName=plumhq&page=1&items=99"
kula = {"meta": {"count": 3, "pages": 1}, "data": [
    {"id": 3920 + i, "title": t, "launch_at": "2026-06-30T11:57:45.000Z",
     "ats_job": {"offices": [{"name": "Plum, 6th Floor, Whitefield", "location": "Bengaluru, Karnataka, India"}]}}
    for i, t in enumerate(["SDET", "Backend Engineer", "Senior Software Engineer"])
]}
kula_found = scrape.find_source([(kula_url, kula)], None, None, page_url=kula_url, url_base=None)
assert kula_found, "Kula ats_job_posts should be read by its known shape"
assert kula_found[2][0]["url"] == "https://careers.kula.ai/plumhq/3920", kula_found[2][0]
assert kula_found[2][0]["location"] == "Bengaluru, Karnataka, India", "the office's location, not its street address"
assert scrape.source_total([(kula_url, kula)], kula_found[0], kula_found[1]) == 3, "meta.count is the total"
pj_url = "https://api.pyjamahr.com/api/career/jobs/?company_uuid=8B92017E1E&page=1&is_careers_page=true"
pj = {"count": 64, "next": None, "results": [
    {"id": 413349 + i, "slug": "x", "title": t, "location": "Bengaluru"}
    for i, t in enumerate(["Backend Engineer", "Data Engineer", "Android Developer"])
]}
pj_found = scrape.find_source([(pj_url, pj)], None, None, page_url=pj_url, url_base=None)
assert pj_found and pj_found[2][0]["url"] == (
    "https://app.pyjamahr.com/careers?company_uuid=8B92017E1E&job_id=413349"
), pj_found
assert scrape.source_total([(pj_url, pj)], pj_found[0], pj_found[1]) == 64, "under-coverage stays visible"
pj_slug_url = "https://api.pyjamahr.com/api/career/jobs/?company_slug=dodo-payments&page=1"
pj_slug = {"count": 3, "results": [dict(r, slug=f"backend-engineer-{i}") for i, r in enumerate(pj["results"])]}
pj_slug_found = scrape.find_source([(pj_slug_url, pj_slug)], None, None, page_url=pj_slug_url, url_base=None)
assert pj_slug_found[2][0]["url"] == "https://jobs.pyjamahr.com/dodo-payments/backend-engineer-0", pj_slug_found[2][0]
# MyNextHire: the posting URL carries the id inside base64 JSON. This exact URL
# was checked by hand to open Amagi's req 436 rather than the list.
mnh_url = "https://amagi.mynexthire.com/employer/careers/reqlist/get"
mnh = {"requesterTitle": "", "reqDetailsBOList": [
    {"reqId": 436 + i, "reqTitle": t, "location": "Bangalore", "approvedOn": "2026-09-16T15:41:52.395+0000"}
    for i, t in enumerate(["Technical Project Lead", "Software Development Engineer III", "Platform Engineer"])
]}
mnh_found = scrape.find_source([(mnh_url, mnh)], None, None, page_url=mnh_url, url_base=None)
assert mnh_found[2][0]["url"] == (
    "https://amagi.mynexthire.com/employer/jobs/careers#?src=careers&p="
    "eyJwYWdlVHlwZSI6ImpkIiwiY3ZTb3VyY2UiOiJjYXJlZXJzIiwicmVxSWQiOjQzNiwicmVxdWVzdGVyIjp7ImlkIjoiIiwiY29kZSI6IiIsIm5hbWUiOiIifSwicGFnZSI6ImNhcmVlcnMiLCJidWZpbHRlciI6LTEsImN1c3RvbUZpZWxkcyI6e319"
    "&page=careers"
), mnh_found[2][0]["url"]
# Dedupe keeps postings that differ only in the query or fragment.
pj_roles = [scrape.to_role({"name": "Kuku FM", "id": "k"}, "xhr", r) for r in pj_found[2]]
pj_roles.append(dict(pj_roles[0], url=pj_roles[0]["url"] + "&utm_source=x"))
kept, _ = scrape.dedupe_scraped(pj_roles, [])
assert len(kept) == 3, f"three job_ids are three roles, and a utm_ copy is not a fourth: {len(kept)}"
mnh_roles = [scrape.to_role({"name": "Amagi", "id": "a"}, "xhr", r) for r in mnh_found[2]]
assert len(scrape.dedupe_scraped(mnh_roles, [])[0]) == 3, "fragment-identified postings stay distinct"
wk_url = "https://apply.workable.com/api/v3/accounts/tiger-analytics/jobs"
wk = {"total": 3, "results": [
    {"shortcode": f"39D0D71D6{i}", "title": t, "published": "2026-09-01T00:00:00.000Z",
     "location": {"country": "India", "city": "Bengaluru", "region": "Karnataka"}}
    for i, t in enumerate(["Data Engineer", "ML Engineer", "Software Engineer"])
]}
wk_found = scrape.find_source([(wk_url, wk)], None, None, page_url=wk_url, url_base=None)
assert wk_found[2][0]["url"] == "https://apply.workable.com/tiger-analytics/j/39D0D71D60/", wk_found[2][0]
assert scrape.find_source(
    [("https://apply.workable.com/api/v3/accounts/tiger-analytics/jobs/filters", wk)], None, None,
    page_url=wk_url, url_base=None,
) is None, "the filters endpoint is not the job list"
wk_india = {"total": 4, "results": wk["results"][:1]}
assert scrape.source_total([(wk_url, {"total": 161, "results": wk["results"]}), (wk_url, wk_india)],
                           wk_found[0], wk_found[1]) == 4, "the refined list's total, not the unfiltered one"

# Experience: read from a board's field, a title, or near "experience" in a JD.
pe = discover.parse_experience
assert pe("2 - 5 Years", structured=True) == (2, 5)
assert pe("3+ yrs", structured=True) == (3, None)
assert pe("Fresher", structured=True) == (0, 0)
assert pe("We need 3+ years of experience building backend services") == (3, None)
assert pe("Experience: 1–3 years in Python") == (1, 3)
assert pe("Minimum of 5 years of professional experience") == (5, None)
assert pe("Founded 12 years ago, we serve 40 countries.") is None, "a company's age is not a requirement"
assert pe("SDE 2 (Backend)") is None, "SDE 2 is a level, not years"
fits = discover.experience_fits
assert fits({"exp_min": 2, "exp_max": 5}, (0, 3)), "overlap: 2-5 is open to someone with 2"
assert fits({"exp_min": 3, "exp_max": None}, (0, 3)), "a 3+ bar is reachable from 0-3"
assert not fits({"exp_min": 6, "exp_max": 9}, (0, 3))
assert not fits({"exp_min": 5, "exp_max": None}, (0, 3))
assert fits({"exp_min": None, "exp_max": None}, (0, 3)), "unstated is kept by default"
assert not fits({"exp_min": None, "exp_max": None}, (0, 3), required=True)
jd_role = discover._role({"name": "Acme"}, "lever", title="Software Engineer", url="u",
                         locations=["Bengaluru"], posted=None,
                         description="About us... Requirements: 6+ years of experience with Go.")
assert (jd_role["exp_min"], jd_role["exp_max"], jd_role["experience"]) == (6, None, "6+ yrs"), jd_role
assert not discover.filter_roles([jd_role], keywords=["software"], experience=(0, 3))
# Board fields: Darwinbox's experience_from/to come through the known map.
assert dbx_found[1]["exp_min"] == "experience_from"
dbx_exp = dict(dbx["data"][0], experience_from="6", experience_to="9")
assert scrape.row_experience(dbx_exp, dbx_found[1]) == (6, 9)
assert scrape.row_experience({"minExperience": 3, "maxExperience": 0}, {"exp_min": "minExperience", "exp_max": "maxExperience"}) == (3, None), "a max of 0 under a min of 3 is unset"
assert scrape.guess_experience_keys([{"min_experience": 1, "max_experience": 3}] * 3) == {"exp_min": "min_experience", "exp_max": "max_experience"}
assert scrape.card_experience(["Bangalore, Karnataka, India", "6 - 9 Years", "Full Time"]) == (6, 9)
assert scrape.experience_range("0-3") == {"min": 0, "max": 3, "required": False}
f = scrape.load_filters({"experience": "0-3"})
senior = scrape.to_role({"name": "Locus", "id": "l"}, "links", {"title": "Staff Engineer", "url": "https://x/1", "location": "Bangalore", "experience": (8, 12)})
young = scrape.to_role({"name": "Locus", "id": "l"}, "links", {"title": "Software Engineer", "url": "https://x/2", "location": "Bangalore", "experience": (1, 3)})
assert [r["title"] for r in scrape.filter_scraped([senior, young], f)] == ["Software Engineer"]

# The per-company snapshot behind the Open roles tab.
acme = {"id": "c-acme", "name": "Acme", "tier": "growth", "category": "startup"}
ok_run = {"company": acme, "problem": None, "detail": "", "strategy": "links",
          "rows": [{"title": "Backend Engineer", "url": "https://acme.example/jobs/1", "location": "Bengaluru"}]}
snap = scrape.merge_scraped({}, [ok_run], as_of="2026-10-05")
assert snap["c-acme"]["read_on"] == "2026-10-05" and len(snap["c-acme"]["roles"]) == 1
assert "description" not in snap["c-acme"]["roles"][0], "the snapshot carries no JD text"
failed = {**ok_run, "problem": "timeout", "detail": "no result", "rows": []}
snap2 = scrape.merge_scraped(snap, [failed], as_of="2026-10-06")
assert snap2["c-acme"]["roles"] and snap2["c-acme"]["read_on"] == "2026-10-05", "a failed read keeps yesterday's roles"
assert snap2["c-acme"]["problem"].startswith("timeout")
empty = {**ok_run, "rows": [], "detail": "the page says there are no openings"}
assert scrape.merge_scraped(snap, [empty], as_of="2026-10-06")["c-acme"]["roles"] == [], "a real zero clears them"
scrape.save_scraped(snap)
assert scrape.load_scraped()["c-acme"]["name"] == "Acme"

import openroles

saved = openroles.save_filters({"experience": {"min": 0, "max": 2}, "exclude_seniority": ["senior"]})
assert saved["experience"] == {"min": 0, "max": 2, "required": False} and saved["exclude_seniority"] == ["senior"]
off = openroles.save_filters({"experience": None})
assert off["experience"] is None, "switching the filter off must not fall back to the 0-3 default"
for bad in ({"experience": "lots"}, {"exclude_seniority": ["wizard"]}, {"keywords": []}, {"locations": ["x"]}):
    try:
        openroles.save_filters(bad)
    except ValueError:
        pass
    else:
        raise AssertionError(f"save_filters accepted {bad}")
openroles.save_filters({"experience": "0-3", "exclude_seniority": ["senior", "intern"]})
on_board = {c["name"] for c in companies.list_companies()}
sugg = openroles.suggestions()
assert sugg and not any(x["name"] in on_board for x in sugg), "suggestions never repeat the board"
first = sugg[0]["name"]
assert openroles.add_suggestions([first])["added"] == 1
assert all(x["name"] != first for x in openroles.suggestions())
assert openroles.scrape_status()["running"] is False
print(f"open roles OK {len(sugg)} suggestions, filters round-trip, snapshot keeps a failed read's roles")
assert not scrape.says_no_openings("Open positions: Backend Engineer, 14 openings in Pune")

# Entry links: a brochure careers page is followed to its list; social links never are.
brochure = [
    {"href": "https://www.linkedin.com/company/acme", "text": "Join us on LinkedIn"},
    {"href": "https://acme.example/about", "text": "About us"},
    {"href": "https://acme.example/careers/life", "text": "Life at Acme"},
    {"href": "https://acme.example/careers/openings", "text": "View open roles"},
]
assert scrape.pick_entry_link(brochure, [], "https://acme.example/careers") == "https://acme.example/careers/openings"
on_ats = brochure + [{"href": "https://acme.wd3.myworkdayjobs.com/External", "text": "Apply"}]
assert scrape.pick_entry_link(on_ats, [], "https://acme.example/careers") == "https://acme.wd3.myworkdayjobs.com/External"
embedded = ["https://boards.greenhouse.io/embed/job_board?for=acme"]
assert scrape.pick_entry_link(brochure, embedded, "https://acme.example/careers") == embedded[0], "an embedded board is the list"
assert scrape.pick_entry_link(brochure[:3], [], "https://acme.example/careers") is None
assert scrape.pick_entry_link(
    [{"href": "https://acme.example/go/Demand/95/", "text": "Explore Job Opportunities"}], [], "https://acme.example/"
) == "https://acme.example/go/Demand/95/", "HCLTech's wording"
assert scrape.pick_entry_link([{"href": "https://acme.example/careers", "text": "Careers"}], [], "https://acme.example/") == "https://acme.example/careers"
assert scrape.site_root("https://acme.example/a/b?c=1") == "https://acme.example/"
assert scrape.listing_url_for("https://g24.darwinbox.in/ms/candidate/main/candidate/login") == \
    "https://g24.darwinbox.in/ms/candidatev2/main/careers/allJobs"
assert scrape.listing_url_for("https://apply.workable.com/tiger-analytics/j/23D676B06C/") == \
    "https://apply.workable.com/tiger-analytics/"
assert scrape.listing_url_for("https://acme.example/careers") == "https://acme.example/careers"

# Plausibility: case studies and a repeated nav item are not job lists.
def _titles(*ts):
    return [{"title": t} for t in ts]
assert scrape.looks_like_jobs(_titles("Senior Category Manager", "Backend Engineer", "GTM recruiter"))
assert not scrape.looks_like_jobs(_titles(
    "Hapi Cloud hit an 85% utilization rate", "Fluxx boosted customer engagement",
    "Gamify slashed implementation timelines"))
assert not scrape.looks_like_jobs(_titles("Application", "Application", "Application"))
assert not scrape.looks_like_jobs([])

# Walls are reported, never bypassed.
assert scrape.classify_page(403, "") == "HTTP 403"
assert scrape.classify_page(200, "Please verify you are human to continue")
assert scrape.classify_page(200, "Open roles: Backend Engineer") is None
assert scrape.classify_exception(RuntimeError("Page.goto: net::ERR_NAME_NOT_RESOLVED")) == "navigation"

# Filtering: a role with no readable location is trusted to the URL's own filter.
filt = {
    "keywords": ["engineer"], "locations": list(discover.INDIA_TOKENS),
    "exclude_keywords": list(discover.NOISE_KEYWORDS), "include_remote": False, "max_age_days": None,
}
co = {"id": "c1", "name": "Globex", "category": "product", "tier": "tier1", "careers_url": "https://g.example/jobs"}
mixed = [
    scrape.to_role(co, "links", {"title": "Backend Engineer", "url": "https://g.example/jobs/1", "location": ""}),
    scrape.to_role(co, "xhr", {"title": "SRE", "url": "https://g.example/jobs/2", "location": "Pune, India"}),
    scrape.to_role(co, "xhr", {"title": "Platform Engineer", "url": "https://g.example/jobs/3", "location": "Austin, US"}),
    scrape.to_role(co, "xhr", {"title": "Sales Engineer", "url": "https://g.example/jobs/4", "location": "Pune, India"}),
]
kept_titles = sorted(r["title"] for r in scrape.filter_scraped(mixed, {**filt, "keywords": ["engineer", "sre"]}))
assert kept_titles == ["Backend Engineer", "SRE"], kept_titles

# Assembly: tier order, dedupe against jobs.json, new-since-last-run, links and failures.
co2 = {"id": "c2", "name": "Initech", "category": "mnc", "tier": "faang", "careers_url": "https://i.example/c"}
co3 = {"id": "c3", "name": "Hooli", "category": "mnc", "tier": "tier2", "careers_url": "https://h.example/c"}
co4 = {"id": "c4", "name": "LinkedIn", "category": "mnc", "tier": "tier1", "careers_url": "https://www.linkedin.com/jobs"}
assert scrape.skip_reason(co4, None), "linkedin is skipped outright"
results = [
    {"company": co, "strategy": "links", "rows": [
        {"title": "Backend Engineer", "url": "https://g.example/jobs/1", "location": ""},
        {"title": "Platform Engineer", "url": "https://g.example/jobs/3", "location": "Austin, US"}],
     "problem": None, "detail": "", "search_url": "https://g.example/jobs?q=software+engineer", "snapshot": None},
    {"company": co2, "strategy": "xhr", "rows": [
        {"title": "Software Engineer", "url": "https://i.example/c/9", "location": "Bengaluru, India"}],
     "problem": None, "detail": "", "search_url": "https://i.example/c?location=India", "snapshot": None},
    {"company": co3, "strategy": None, "rows": [], "problem": "blocked", "detail": "HTTP 403",
     "search_url": "https://h.example/c", "snapshot": "data/scrape-snapshots/x/hooli"},
    {"company": co4, "strategy": None, "rows": [], "problem": "skipped",
     "detail": "login wall and terms of service", "search_url": "https://www.linkedin.com/jobs", "snapshot": None},
]
built = scrape.assemble(results, filt, [], {"https://i.example/c/9"})
assert [r["company"] for r in built["roles"]] == ["Initech", "Globex"], "faang sorts before tier1"
assert [r["is_new"] for r in built["roles"]] == [False, True], "an earlier export's link is not new"
assert built["totals"]["ok"] == 2 and built["totals"]["failed"] == 1 and built["totals"]["skipped"] == 1
assert [f["company"] for f in built["failures"]] == ["Hooli"], "a skip is not a failure"
assert {a["company"] for a in built["attention"]} == {"Hooli", "LinkedIn"}, "the sheet names every unread company"
assert len(built["links"]) == 4, "every company still gets its filtered link"
assert any(l["status"] == "ok" and l["found"] == 1 for l in built["links"] if l["company"] == "Initech")
tracked_url = [{"url": "https://i.example/c/9", "organisation": "", "job_title": ""}]
assert scrape.assemble(results, filt, tracked_url, set())["totals"]["already_tracked"] == 1
# Same title, different link: both are real openings. A title match only counts for a
# tracked record that has no URL of its own.
twins = [
    {"company": co, "strategy": "xhr", "problem": None, "detail": "", "snapshot": None,
     "search_url": "https://g.example/jobs",
     "rows": [{"title": "Software Engineer", "url": f"https://g.example/jobs/{n}", "location": "Pune, India"}
              for n in (1, 2, 3)]},
]
assert len(scrape.assemble(twins, filt, [], set())["roles"]) == 3, "same-title postings are distinct roles"
by_hand = [{"url": None, "organisation": "Globex", "job_title": "software engineer"}]
assert scrape.assemble(twins, filt, by_hand, set())["totals"]["already_tracked"] == 3
with_url = [{"url": "https://elsewhere/1", "organisation": "Globex", "job_title": "Software Engineer"}]
assert scrape.assemble(twins, filt, with_url, set())["totals"]["already_tracked"] == 0

# Config merge: a failure keeps the working config; layout_changed keeps last_count.
saved = {"company_id": "c3", "name": "Hooli", "strategy": "xhr", "last_count": 40, "endpoint": "e"}
bad = scrape.merge_site(saved, co3, results[2], "https://h.example/c")
assert bad["strategy"] == "xhr" and bad["endpoint"] == "e" and bad["last_error"].startswith("blocked")
shrunk = {"company": co3, "strategy": "xhr", "rows": [1], "learned": {}, "problem": "layout_changed", "detail": "1 vs 40"}
assert scrape.merge_site(saved, co3, shrunk, "u")["last_count"] == 40
healed = {"company": co3, "strategy": "links", "rows": [1, 2, 3], "learned": {}, "problem": None, "detail": ""}
assert scrape.merge_site(saved, co3, healed, "u")["last_count"] == 3
under = {"company": co3, "strategy": "xhr", "rows": [1, 2], "learned": {"endpoint": "e2"},
         "problem": "partial", "detail": "read 2 of 244 roles"}
merged_under = scrape.merge_site(saved, co3, under, "u")
assert merged_under["last_count"] == 2 and merged_under["endpoint"] == "e2"
assert merged_under["last_error"].startswith("partial"), "a partial read keeps its config but says so"

# Selection: only companies with no JSON board and a link; flags narrow it.
board_rows = [
    {**co, "ats_provider": "none", "status": "researching"},
    {**co2, "ats_provider": "none", "status": "target"},
    {**co3, "ats_provider": "greenhouse", "status": "researching"},
    {"id": "c5", "name": "Dead", "tier": "tier1", "ats_provider": "none", "status": "researching", "careers_url": None},
    {"id": "c6", "name": "Meh", "tier": "tier1", "ats_provider": "none", "status": "not_interested",
     "careers_url": "https://m.example"},
]
assert {c["name"] for c in scrape.select(board_rows, tiers=[], statuses=[], names=[])} == {"Globex", "Initech"}
assert [c["name"] for c in scrape.select(board_rows, tiers=["tier1"], statuses=["target"], names=[])] == ["Globex", "Initech"]
assert [c["name"] for c in scrape.select(board_rows, tiers=[], statuses=[], names=["meh"])] == ["Meh"]

# A followed entry link becomes the site's search_url.
hopped = {"company": co3, "strategy": "links", "rows": [1], "problem": None, "detail": "",
          "learned": {"search_url": "https://h.example/careers/openings"}}
assert scrape.merge_site(saved, co3, hopped, "https://h.example/c")["search_url"] == "https://h.example/careers/openings"

# Manual review: a history per company, latest wins; the queue offers readable sites first.
reviewed_site = scrape.record_review(None, co, "needs_fix", "missed the Pune roles", on="2026-10-05")
reviewed_site = scrape.record_review(reviewed_site, co, "verified", on="2026-10-06")
assert len(reviewed_site["reviews"]) == 2 and scrape.review_label(reviewed_site) == "verified 2026-10-06"
try:
    scrape.record_review(None, co, "looks fine")
    raise AssertionError("an unknown verdict must be rejected")
except ValueError:
    pass
pool_rows = [
    {**co, "tier": "tier2"},                      # reviewed
    {**co2, "tier": "faang"},                     # pending, readable
    {**co4, "tier": "tier1"},                     # pending, but LinkedIn is skipped
    {"id": "c7", "name": "Umbrella", "tier": "tier1", "careers_url": "https://u.example/jobs"},
]
queue_sites = {"c1": reviewed_site}
assert [c["name"] for c in scrape.review_queue(pool_rows, queue_sites)] == ["Initech", "Umbrella"]
assert [c["name"] for c in scrape.review_queue(pool_rows, queue_sites, limit=5)][-1] == "LinkedIn"

# Needs-your-help: a reason is required, and the list reads it back by tier.
try:
    scrape.flag_help(None, co, "  ")
    raise AssertionError("a help flag without a reason must be rejected")
except ValueError:
    pass
queue_sites["c7"] = scrape.flag_help(None, pool_rows[3], "404, needs a new careers link", on="2026-10-05")
assert [(h["company"], h["why"]) for h in scrape.help_rows(pool_rows, queue_sites)] == [
    ("Umbrella", "404, needs a new careers link")]
assert scrape.selection_slug(["tier1", "faang"], ["target"], []) == "faang-tier1-target"
assert scrape.selection_slug([], [], ["D. E. Shaw India"]) == "d-e-shaw-india"

# Filters file: created on first use, null expands to the India list, flags override.
loaded = scrape.load_filters({"keywords": ["rust"], "include_remote": True})
assert (jsonstore.DATA_DIR / scrape.FILTERS_FILE).exists()
assert loaded["keywords"] == ["rust"] and loaded["include_remote"] is True
assert "bengaluru" in loaded["locations"] and loaded["exclude_keywords"] == list(discover.NOISE_KEYWORDS)
assert scrape.load_filters()["keywords"] == scrape.DEFAULT_FILTERS["keywords"], "overrides are per-run"

# Workbook: four sheets, real hyperlinks, scraped text pinned to strings, new-since diff.
from openpyxl import load_workbook  # noqa: E402

xl_dir = Path(tmp) / "exports"
hostile = {**built["roles"][0], "title": "=HYPERLINK(\"http://evil\",\"x\")"}
xl = roles_xlsx.write(
    xl_dir / "open-roles-2026-10-04.xlsx", [hostile, *built["roles"][1:]],
    [{**built["links"][0], "reviewed": "verified 2026-10-04"}, *built["links"][1:]],
    built["failures"], [("Run date", "2026-10-04")],
    scrape.help_rows(pool_rows, queue_sites),
)
wb = load_workbook(xl)
assert wb.sheetnames == [roles_xlsx.ROLES_SHEET, roles_xlsx.LINKS_SHEET, roles_xlsx.FAILURES_SHEET,
                         roles_xlsx.HELP_SHEET, roles_xlsx.SUMMARY_SHEET]
assert wb[roles_xlsx.HELP_SHEET]["A2"].value == "Umbrella" and wb[roles_xlsx.HELP_SHEET]["C2"].hyperlink
assert wb[roles_xlsx.LINKS_SHEET]["F2"].value == "verified 2026-10-04"
# A partial run is compared only with earlier runs of the same selection.
roles_xlsx.write(xl_dir / "open-roles-2026-10-03-faang.xlsx", [], [], [], [])
assert roles_xlsx.latest_before(xl_dir, xl_dir / "open-roles-2026-10-05-faang.xlsx").name == "open-roles-2026-10-03-faang.xlsx"
assert roles_xlsx.latest_before(xl_dir, xl_dir / "open-roles-2026-10-05.xlsx") == xl
cell = wb[roles_xlsx.ROLES_SHEET]["C2"]
assert cell.data_type == "s" and cell.value.startswith("="), "a scraped '=...' title must not become a formula"
assert wb[roles_xlsx.ROLES_SHEET]["H2"].hyperlink.target == built["roles"][0]["url"]
assert wb[roles_xlsx.ROLES_SHEET].freeze_panes == "A2"
assert wb[roles_xlsx.FAILURES_SHEET]["A2"].value == "Hooli"
assert roles_xlsx.read_previous_urls(xl) == {r["url"] for r in built["roles"]}
assert roles_xlsx.latest_before(xl_dir, xl_dir / "open-roles-2026-10-05.xlsx") == xl
assert roles_xlsx.latest_before(xl_dir, xl) is None
assert roles_xlsx.read_previous_urls(xl_dir / "missing.xlsx") == set()
print(f"scrape OK   {len(built['roles'])} roles assembled, {len(rows)} mapped from a workday-shaped payload")

# ------------------------------------------------------------------ prompt
import ai  # noqa: E402

prompt = ai.system_prompt()
# The prompt carries LaTeX macro examples; an f-string would eat the braces and
# every chat turn would die before reaching the model.
for literal in (
    r"\resumeSubheading{org}{right}{role}{date}",
    r"\pdfgentounicode=1",
    r"O(n \log n)",
):
    assert literal in prompt, f"mangled in prompt: {literal}"
names = {tool.name for tool in ai.TOOLS}
import re  # noqa: E402

for candidate in set(
    re.findall(r"\b(?:add|get|list|log|update|save|link|set|flag|resolve|cap)_[a-z_]+\b", prompt)
):
    assert candidate in names, f"prompt references a missing tool: {candidate}"
assert ai.build_options(None).system_prompt == prompt
# The prompt is a plain literal, so the cap it states can drift from the constant.
assert f"gets {dsa.ATTEMPT_CAP_MINUTES} minutes on the clock" in prompt, "prompt and cap disagree"

# The cap is enforced by the server putting the clock in front of the turn, not
# by the model remembering to look.
assert ai.with_clock_note("hello") == "hello"
overrun = dsa.create_problem({"title": "Overrun"})
dsa.set_timer(overrun["id"], "start")
_rewind(overrun["id"], (dsa.ATTEMPT_CAP_MINUTES + 1) * 60)
noted = ai.with_clock_note("one more minute")
assert noted.startswith("[stopwatch] Overrun is at 41 min"), noted
assert overrun["id"] in noted and noted.endswith("\n\none more minute"), noted
dsa.cap_attempt(overrun["id"])
assert ai.clock_note() == ""
print(f"prompt OK   {len(names)} tools, {len(prompt)} chars")

# The Windows loop bridge: uvicorn --reload hands us a SelectorEventLoop, which
# cannot spawn the CLI subprocess. Losing this bridge breaks every chat turn
# while leaving the rest of the app working, so it is easy to miss.
assert hasattr(ai, "_needs_proactor_bridge"), "Windows event-loop bridge is gone"
assert not ai._needs_proactor_bridge(), "no running loop should mean no bridge"


async def _is_bridged() -> bool:
    return ai._needs_proactor_bridge()


if sys.platform == "win32":
    selector_loop = asyncio.SelectorEventLoop()
    try:
        assert selector_loop.run_until_complete(_is_bridged()), "selector loop must bridge"
    finally:
        selector_loop.close()
    proactor_loop = asyncio.ProactorEventLoop()
    try:
        assert not proactor_loop.run_until_complete(_is_bridged()), "proactor needs no bridge"
    finally:
        proactor_loop.close()
    print("loop bridge OK  selector->bridged, proactor->direct")

print("\nALL CHECKS PASSED")
